"""깨짐없는 PPT 변환 엔진 — Vercel 서버리스 / 로컬 공용 순수 모듈.

핵심: PDF에서 추출한 한글 텍스트를 PPTX에 쓸 때 OOXML의
동아시아(ea) 서체를 명시적으로 지정해 'ㅁㅁㅁ' 깨짐을 원천 차단한다.
"""
import io
from http.server import BaseHTTPRequestHandler

import pymupdf
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.dml.color import RGBColor

FONTS = {
    "malgun": "맑은 고딕",
    "noto": "Noto Sans KR",
    "pretendard": "Pretendard",
}
ASPECTS = {"16:9": (13.333, 7.5), "4:3": (10, 7.5)}
EMU_PER_PT = 12700  # 914400 EMU/inch / 72 pt/inch
MIN_CHARS_PER_PAGE = 30  # 이보다 적으면 이미지 PDF로 판단
MAX_PAGES = 80


class ConvertError(Exception):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code
        self.message = message


ERRORS = {
    "not_pdf": "PDF 파일이 아니에요. .pdf 파일을 올려주세요.",
    "broken": "PDF 파일을 열 수 없어요. 파일이 손상된 것 같아요.",
    "encrypted": "암호가 걸린 PDF예요. 암호를 해제한 뒤 올려주세요.",
    "empty": "빈 PDF예요. 페이지가 하나도 없어요.",
    "too_many": "페이지가 너무 많아요. 80페이지 이하로 나눠서 올려주세요.",
}


def set_korean_font(run, font_name):
    """텍스트 run에 한글(ea) 서체를 명시 — 깨짐 방지 핵심."""
    run.font.name = font_name  # latin 서체
    rPr = run._r.get_or_add_rPr()
    for ea in list(rPr.findall(qn("a:ea"))):
        rPr.remove(ea)
    ea = parse_xml('<a:ea %s typeface="%s"/>' % (nsdecls("a"), font_name))
    latin = rPr.find(qn("a:latin"))
    if latin is not None:
        latin.addnext(ea)
    else:
        rPr.append(ea)
    rPr.set("lang", "ko-KR")


def _open_doc(pdf_bytes):
    if not pdf_bytes[:5] == b"%PDF-":
        raise ConvertError("not_pdf", ERRORS["not_pdf"])
    try:
        doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    except Exception:
        raise ConvertError("broken", ERRORS["broken"])
    if doc.needs_pass:
        raise ConvertError("encrypted", ERRORS["encrypted"])
    if doc.page_count == 0:
        raise ConvertError("empty", ERRORS["empty"])
    if doc.page_count > MAX_PAGES:
        raise ConvertError("too_many", ERRORS["too_many"])
    return doc


def diagnose(pdf_bytes):
    """PDF를 분석해 변환 가능 여부와 권장 모드를 반환."""
    doc = _open_doc(pdf_bytes)
    pages = []
    total_chars = 0
    for i, page in enumerate(doc):
        try:
            text = page.get_text().strip()
        except Exception:
            text = ""
        try:
            n_images = len(page.get_images())
        except Exception:
            n_images = 0
        pages.append({"page": i + 1, "chars": len(text), "images": n_images})
        total_chars += len(text)
    avg = total_chars / max(1, len(pages))
    extractable = avg >= MIN_CHARS_PER_PAGE
    return {
        "pages": len(pages),
        "total_chars": total_chars,
        "avg_chars": round(avg, 1),
        "extractable": extractable,
        "mode": "fast" if extractable else "needs_ocr",
        "detail": pages,
    }


def _extract_png(doc, xref):
    try:
        pix = pymupdf.Pixmap(doc, xref)
        if pix.n - pix.alpha > 3:  # CMYK → RGB
            pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
        return pix.tobytes("png")
    except Exception:
        return None


def convert_pdf_to_pptx(pdf_bytes, font_key="malgun", aspect="16:9"):
    """고속 모드: PDF 텍스트/이미지를 직접 추출해 PPTX 생성.
    Returns (pptx_bytes | None, report dict)."""
    diag = diagnose(pdf_bytes)
    font_name = FONTS.get(font_key, FONTS["malgun"])
    if not diag["extractable"]:
        return None, {
            "mode": "needs_ocr",
            "diagnosis": diag,
            "message": "텍스트 추출이 어려운 PDF예요. 정밀 모드(Gemini)로 변환할 수 있어요.",
        }

    doc = _open_doc(pdf_bytes)
    prs = Presentation()
    sw_in, sh_in = ASPECTS.get(aspect, ASPECTS["16:9"])
    prs.slide_width = Emu(int(sw_in * 914400))
    prs.slide_height = Emu(int(sh_in * 914400))
    blank = prs.slide_layouts[6]
    warnings = []

    for page in doc:
        slide = prs.slides.add_slide(blank)
        pw, ph = page.rect.width, page.rect.height
        if pw <= 0 or ph <= 0:
            continue
        scale = min(prs.slide_width / (pw * EMU_PER_PT),
                    prs.slide_height / (ph * EMU_PER_PT))
        ox = (prs.slide_width - pw * EMU_PER_PT * scale) / 2
        oy = (prs.slide_height - ph * EMU_PER_PT * scale) / 2

        def X(x): return Emu(int(ox + x * EMU_PER_PT * scale))
        def Y(y): return Emu(int(oy + y * EMU_PER_PT * scale))
        def W(w): return Emu(max(1, int(w * EMU_PER_PT * scale)))

        # 이미지 먼저 (텍스트 아래 깔림)
        try:
            for img in page.get_images(full=True):
                xref = img[0]
                png = _extract_png(doc, xref)
                if not png:
                    continue
                try:
                    rects = page.get_image_rects(xref)
                except Exception:
                    continue
                for rect in rects:
                    try:
                        slide.shapes.add_picture(
                            io.BytesIO(png), X(rect.x0), Y(rect.y0),
                            W(rect.width), W(rect.height))
                    except Exception:
                        continue
        except Exception:
            pass

        # 텍스트 블록 → 텍스트 상자
        try:
            d = page.get_text("dict")
        except Exception:
            d = {"blocks": []}
        for block in d.get("blocks", []):
            if block.get("type", 0) != 0:
                continue
            x0, y0, x1, y1 = block["bbox"]
            if x1 - x0 < 2 or y1 - y0 < 2:
                continue
            try:
                txBox = slide.shapes.add_textbox(X(x0), Y(y0), W(x1 - x0), W(y1 - y0))
            except Exception:
                continue
            tf = txBox.text_frame
            tf.word_wrap = True
            first = True
            for line in block.get("lines", []):
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                p.alignment = PP_ALIGN.LEFT
                for span in line.get("spans", []):
                    text = span.get("text", "").replace(" ", " ")
                    if not text:
                        continue
                    run = p.add_run()
                    run.text = text
                    size = span.get("size", 12) * scale
                    run.font.size = Pt(max(6, min(72, size)))
                    set_korean_font(run, font_name)
                    flags = span.get("flags", 0)
                    run.font.bold = bool(flags & 16)
                    run.font.italic = bool(flags & 2)
                    color = span.get("color", 0)
                    if color:
                        try:
                            run.font.color.rgb = RGBColor(
                                (color >> 16) & 255, (color >> 8) & 255, color & 255)
                        except Exception:
                            pass

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), {
        "mode": "fast",
        "pages": doc.page_count,
        "font": font_name,
        "aspect": aspect,
        "diagnosis": diag,
        "warnings": warnings,
    }


def build_from_json(slides, font_key="malgun", aspect="16:9"):
    """정밀 모드: Gemini가 추출한 슬라이드 JSON → PPTX.
    slides = [{"title": str, "bullets": [str], "notes": str?}]"""
    if not isinstance(slides, list) or not slides:
        raise ConvertError("bad_json", "슬라이드 데이터가 비어 있어요.")
    font_name = FONTS.get(font_key, FONTS["malgun"])
    prs = Presentation()
    sw_in, sh_in = ASPECTS.get(aspect, ASPECTS["16:9"])
    prs.slide_width = Emu(int(sw_in * 914400))
    prs.slide_height = Emu(int(sh_in * 914400))
    blank = prs.slide_layouts[6]

    for item in slides:
        if not isinstance(item, dict):
            continue
        title = str(item.get("title", "")).strip()
        bullets = item.get("bullets", []) or []
        notes = str(item.get("notes", "")).strip()
        slide = prs.slides.add_slide(blank)

        # 제목
        tbox = slide.shapes.add_textbox(Emu(int(0.7 * 914400)), Emu(int(0.4 * 914400)),
                                        Emu(int((sw_in - 1.4) * 914400)), Emu(int(1.2 * 914400)))
        tf = tbox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = title or "(제목 없음)"
        run.font.size = Pt(30)
        run.font.bold = True
        set_korean_font(run, font_name)
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

        # 강조선
        from pptx.util import Inches
        line = slide.shapes.add_shape(
            1, Emu(int(0.7 * 914400)), Emu(int(1.55 * 914400)),
            Emu(int(1.2 * 914400)), Emu(int(0.06 * 914400)))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(0xFF, 0xD1, 0x66)
        line.line.fill.background()

        # 불릿
        cbox = slide.shapes.add_textbox(Emu(int(0.7 * 914400)), Emu(int(1.9 * 914400)),
                                        Emu(int((sw_in - 1.4) * 914400)), Emu(int((sh_in - 2.6) * 914400)))
        ctf = cbox.text_frame
        ctf.word_wrap = True
        first = True
        for b in bullets[:8]:
            b = str(b).strip()
            if not b:
                continue
            bp = ctf.paragraphs[0] if first else ctf.add_paragraph()
            first = False
            bp.space_after = Pt(8)
            r = bp.add_run()
            r.text = "•  " + b
            r.font.size = Pt(18)
            set_korean_font(r, font_name)
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        if notes:
            slide.notes_slide.placeholders[1].text = notes

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), {
        "mode": "precise",
        "pages": len(slides),
        "font": font_name,
        "aspect": aspect,
    }


# Vercel이 api/*.py를 전부 엔드포인트로 취급하므로,
# 이 모듈 직접 호출 시 404 안내를 반환한다.
class handler(BaseHTTPRequestHandler):
    def _send(self, code, obj):
        import json
        body = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        self._send(404, {"ok": False, "error": "이 주소는 직접 호출할 수 없어요."})

    def do_POST(self):
        self._send(404, {"ok": False, "error": "이 주소는 직접 호출할 수 없어요."})
