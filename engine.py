"""오타잡이 PPT — Vercel 서버리스 / 로컬 공용 순수 모듈.

핵심: NotebookLM 슬라이드 PDF·이미지를 '직접 오타를 고칠 수 있는'
편집 가능 PPTX로 변환한다. 한글 텍스트는 OOXML의 동아시아(ea) 서체를
명시적으로 지정해 'ㅁㅁㅁ' 깨짐을 원천 차단한다.
"""
import io

import pymupdf
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.dml.color import RGBColor

FONTS = {
    "malgun": "맑은 고딕",
    "noto": "Noto Sans KR",
    "pretendard": "Pretendard",
}
ASPECTS = {"16:9": (13.333, 7.5), "4:3": (10, 7.5)}
EMU_PER_PT = 12700  # 914400 EMU/inch / 72 pt/inch
MIN_CHARS_PER_PAGE = 30  # 이보다 적으면 이미지 PDF로 판단
MAX_PAGES = 80


class ConvertError(Exception):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code
        self.message = message


ERRORS = {
    "not_pdf": "PDF 파일이 아니에요. .pdf 파일을 올려주세요.",
    "broken": "PDF 파일을 열 수 없어요. 파일이 손상된 것 같아요.",
    "encrypted": "암호가 걸린 PDF예요. 암호를 해제한 뒤 올려주세요.",
    "empty": "빈 PDF예요. 페이지가 하나도 없어요.",
    "too_many": "페이지가 너무 많아요. 80페이지 이하로 나눠서 올려주세요.",
}


def set_korean_font(run, font_name):
    """텍스트 run에 한글(ea) 서체를 명시 — 깨짐 방지 핵심."""
    run.font.name = font_name  # latin 서체
    rPr = run._r.get_or_add_rPr()
    for ea in list(rPr.findall(qn("a:ea"))):
        rPr.remove(ea)
    ea = parse_xml('<a:ea %s typeface="%s"/>' % (nsdecls("a"), font_name))
    latin = rPr.find(qn("a:latin"))
    if latin is not None:
        latin.addnext(ea)
    else:
        rPr.append(ea)
    rPr.set("lang", "ko-KR")


def _open_doc(pdf_bytes):
    if not pdf_bytes[:5] == b"%PDF-":
        raise ConvertError("not_pdf", ERRORS["not_pdf"])
    try:
        doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    except Exception:
        raise ConvertError("broken", ERRORS["broken"])
    if doc.needs_pass:
        raise ConvertError("encrypted", ERRORS["encrypted"])
    if doc.page_count == 0:
        raise ConvertError("empty", ERRORS["empty"])
    if doc.page_count > MAX_PAGES:
        raise ConvertError("too_many", ERRORS["too_many"])
    return doc


def diagnose(pdf_bytes):
    """PDF를 분석해 변환 가능 여부와 권장 모드를 반환."""
    doc = _open_doc(pdf_bytes)
    pages = []
    total_chars = 0
    for i, page in enumerate(doc):
        try:
            text = page.get_text().strip()
        except Exception:
            text = ""
        try:
            n_images = len(page.get_images())
        except Exception:
            n_images = 0
        pages.append({"page": i + 1, "chars": len(text), "images": n_images})
        total_chars += len(text)
    avg = total_chars / max(1, len(pages))
    extractable = avg >= MIN_CHARS_PER_PAGE
    return {
        "pages": len(pages),
        "total_chars": total_chars,
        "avg_chars": round(avg, 1),
        "extractable": extractable,
        "mode": "fast" if extractable else "needs_ocr",
        "detail": pages,
    }


def _extract_png(doc, xref):
    try:
        pix = pymupdf.Pixmap(doc, xref)
        if pix.n - pix.alpha > 3:  # CMYK → RGB
            pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
        return pix.tobytes("png")
    except Exception:
        return None


def convert_pdf_to_pptx(pdf_bytes, font_key="malgun", aspect="16:9"):
    """고속 모드: PDF 텍스트/이미지를 직접 추출해 PPTX 생성.
    Returns (pptx_bytes | None, report dict)."""
    diag = diagnose(pdf_bytes)
    font_name = FONTS.get(font_key, FONTS["malgun"])
    if not diag["extractable"]:
        return None, {
            "mode": "needs_ocr",
            "diagnosis": diag,
            "message": "텍스트 추출이 어려운 PDF예요. 정밀 모드(Gemini)로 변환할 수 있어요.",
        }

    doc = _open_doc(pdf_bytes)
    prs = Presentation()
    sw_in, sh_in = ASPECTS.get(aspect, ASPECTS["16:9"])
    prs.slide_width = Emu(int(sw_in * 914400))
    prs.slide_height = Emu(int(sh_in * 914400))
    blank = prs.slide_layouts[6]
    warnings = []

    for page in doc:
        slide = prs.slides.add_slide(blank)
        pw, ph = page.rect.width, page.rect.height
        if pw <= 0 or ph <= 0:
            continue
        scale = min(prs.slide_width / (pw * EMU_PER_PT),
                    prs.slide_height / (ph * EMU_PER_PT))
        ox = (prs.slide_width - pw * EMU_PER_PT * scale) / 2
        oy = (prs.slide_height - ph * EMU_PER_PT * scale) / 2

        def X(x): return Emu(int(ox + x * EMU_PER_PT * scale))
        def Y(y): return Emu(int(oy + y * EMU_PER_PT * scale))
        def W(w): return Emu(max(1, int(w * EMU_PER_PT * scale)))

        # 이미지 먼저 (텍스트 아래 깔림)
        try:
            for img in page.get_images(full=True):
                xref = img[0]
                png = _extract_png(doc, xref)
                if not png:
                    continue
                try:
                    rects = page.get_image_rects(xref)
                except Exception:
                    continue
                for rect in rects:
                    try:
                        slide.shapes.add_picture(
                            io.BytesIO(png), X(rect.x0), Y(rect.y0),
                            W(rect.width), W(rect.height))
                    except Exception:
                        continue
        except Exception:
            pass

        # 텍스트 블록 → 텍스트 상자
        try:
            d = page.get_text("dict")
        except Exception:
            d = {"blocks": []}
        for block in d.get("blocks", []):
            if block.get("type", 0) != 0:
                continue
            x0, y0, x1, y1 = block["bbox"]
            if x1 - x0 < 2 or y1 - y0 < 2:
                continue
            try:
                txBox = slide.shapes.add_textbox(X(x0), Y(y0), W(x1 - x0), W(y1 - y0))
            except Exception:
                continue
            tf = txBox.text_frame
            tf.word_wrap = True
            first = True
            for line in block.get("lines", []):
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                p.alignment = PP_ALIGN.LEFT
                for span in line.get("spans", []):
                    text = span.get("text", "").replace(" ", " ")
                    if not text:
                        continue
                    run = p.add_run()
                    run.text = text
                    size = span.get("size", 12) * scale
                    run.font.size = Pt(max(6, min(72, size)))
                    set_korean_font(run, font_name)
                    flags = span.get("flags", 0)
                    run.font.bold = bool(flags & 16)
                    run.font.italic = bool(flags & 2)
                    color = span.get("color", 0)
                    if color:
                        try:
                            run.font.color.rgb = RGBColor(
                                (color >> 16) & 255, (color >> 8) & 255, color & 255)
                        except Exception:
                            pass

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), {
        "mode": "fast",
        "pages": doc.page_count,
        "font": font_name,
        "aspect": aspect,
        "diagnosis": diag,
        "warnings": warnings,
    }


def build_from_json(slides, font_key="malgun", aspect="16:9"):
    """정밀 모드: Gemini가 추출한 슬라이드 JSON → PPTX.
    slides = [{"title": str, "bullets": [str], "notes": str?}]"""
    if not isinstance(slides, list) or not slides:
        raise ConvertError("bad_json", "슬라이드 데이터가 비어 있어요.")
    font_name = FONTS.get(font_key, FONTS["malgun"])
    prs = Presentation()
    sw_in, sh_in = ASPECTS.get(aspect, ASPECTS["16:9"])
    prs.slide_width = Emu(int(sw_in * 914400))
    prs.slide_height = Emu(int(sh_in * 914400))
    blank = prs.slide_layouts[6]

    for item in slides:
        if not isinstance(item, dict):
            continue
        title = str(item.get("title", "")).strip()
        bullets = item.get("bullets", []) or []
        notes = str(item.get("notes", "")).strip()
        slide = prs.slides.add_slide(blank)

        # 제목
        tbox = slide.shapes.add_textbox(Emu(int(0.7 * 914400)), Emu(int(0.4 * 914400)),
                                        Emu(int((sw_in - 1.4) * 914400)), Emu(int(1.2 * 914400)))
        tf = tbox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = title or "(제목 없음)"
        run.font.size = Pt(30)
        run.font.bold = True
        set_korean_font(run, font_name)
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

        # 강조선
        from pptx.util import Inches
        line = slide.shapes.add_shape(
            1, Emu(int(0.7 * 914400)), Emu(int(1.55 * 914400)),
            Emu(int(1.2 * 914400)), Emu(int(0.06 * 914400)))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(0xFF, 0xD1, 0x66)
        line.line.fill.background()

        # 불릿
        cbox = slide.shapes.add_textbox(Emu(int(0.7 * 914400)), Emu(int(1.9 * 914400)),
                                        Emu(int((sw_in - 1.4) * 914400)), Emu(int((sh_in - 2.6) * 914400)))
        ctf = cbox.text_frame
        ctf.word_wrap = True
        first = True
        for b in bullets[:8]:
            b = str(b).strip()
            if not b:
                continue
            bp = ctf.paragraphs[0] if first else ctf.add_paragraph()
            first = False
            bp.space_after = Pt(8)
            r = bp.add_run()
            r.text = "•  " + b
            r.font.size = Pt(18)
            set_korean_font(r, font_name)
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        if notes:
            slide.notes_slide.placeholders[1].text = notes

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), {
        "mode": "precise",
        "pages": len(slides),
        "font": font_name,
        "aspect": aspect,
    }


def build_from_images(images, aspect="16:9"):
    """이미지들을 각 슬라이드에 비율 유지로 가득 채워 PPTX 생성 (빠른 담기용).
    images = [(bytes, name), ...] — JPG/PNG/WebP 모두 PNG로 정규화."""
    from PIL import Image
    if not images:
        raise ConvertError("bad_images", "이미지가 비어 있어요.")
    if len(images) > 20:
        raise ConvertError("too_many", "이미지는 최대 20장까지 올릴 수 있어요.")
    prs = Presentation()
    sw_in, sh_in = ASPECTS.get(aspect, ASPECTS["16:9"])
    sw, sh = Emu(int(sw_in * 914400)), Emu(int(sh_in * 914400))
    prs.slide_width, prs.slide_height = sw, sh
    blank = prs.slide_layouts[6]
    for data, name in images:
        try:
            im = Image.open(io.BytesIO(data))
        except Exception:
            raise ConvertError("bad_images", "이미지를 읽을 수 없어요: %s" % (name or ""))
        if im.mode in ("RGBA", "LA"):
            bg = Image.new("RGB", im.size, (255, 255, 255))
            bg.paste(im, mask=im.split()[-1])
            im = bg
        elif im.mode != "RGB":
            im = im.convert("RGB")
        buf = io.BytesIO()
        im.save(buf, format="PNG")
        png = buf.getvalue()
        iw, ih = im.size
        if iw <= 0 or ih <= 0:
            continue
        slide = prs.slides.add_slide(blank)
        scale = min(sw / iw, sh / ih)
        w, h = Emu(int(iw * scale)), Emu(int(ih * scale))
        left, top = Emu(int((sw - w) / 2)), Emu(int((sh - h) / 2))
        slide.shapes.add_picture(io.BytesIO(png), left, top, w, h)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), {
        "mode": "image",
        "pages": len(images),
        "aspect": aspect,
    }


def render_page_png(pdf_bytes, page_no, width_px=1500):
    """PDF의 특정 페이지를 PNG 바이트로 렌더 (정밀 모드 이미지 크롭용)."""
    doc = _open_doc(pdf_bytes)
    if page_no < 1 or page_no > doc.page_count:
        raise ConvertError("bad_page", "페이지 번호가 범위를 벗어났어요.")
    page = doc[page_no - 1]
    pw = page.rect.width
    if pw <= 0:
        raise ConvertError("bad_page", "페이지를 렌더할 수 없어요.")
    zoom = width_px / pw
    pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom))
    return pix.tobytes("png")


_ALIGN_MAP = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_MAX_ELEMENTS = 80


def _hex_to_rgb(hexstr, default=(0x33, 0x33, 0x33)):
    try:
        h = str(hexstr).strip().lstrip("#")
        if len(h) == 6:
            return (int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
    except Exception:
        pass
    return default


def build_from_layout(slides, sources, font_key="malgun", aspect="16:9"):
    """정밀 모드(레이아웃 재현): Gemini가 추출한 위치 요소 JSON → PPTX.

    slides = [{"page": 1, "notes": str?, "elements": [
        {"type": "text", "x": 0~100, "y": 0~100, "w":.., "h":..,
         "text": str, "font_size": pt, "bold": bool, "italic": bool,
         "color": "RRGGBB", "align": "left|center|right"},
        {"type": "shape", "shape": "rect|ellipse",
         "x":..,"y":..,"w":..,"h":.., "fill": "RRGGBB"},
        {"type": "image", "x":..,"y":..,"w":..,"h":..},
    ]}]
    sources = [png_bytes...] — 슬라이드 순서대로 원본 페이지 이미지 ("image" 요소 크롭용).
    elements는 받은 순서대로 그려 배경→전경 레이어를 재현한다.
    오타는 사용자가 직접 고치므로 텍스트를 절대 임의로 수정하지 않는다.
    """
    if not isinstance(slides, list) or not slides:
        raise ConvertError("bad_json", "슬라이드 데이터가 비어 있어요.")
    font_name = FONTS.get(font_key, FONTS["malgun"])
    prs = Presentation()
    sw_in, sh_in = ASPECTS.get(aspect, ASPECTS["16:9"])
    sw, sh = Emu(int(sw_in * 914400)), Emu(int(sh_in * 914400))
    prs.slide_width, prs.slide_height = sw, sh
    blank = prs.slide_layouts[6]

    from PIL import Image as PILImage

    def NX(x):
        return Emu(int(max(0, min(100, float(x or 0))) / 100 * sw))
    def NY(y):
        return Emu(int(max(0, min(100, float(y or 0))) / 100 * sh))
    def NW(w):
        return Emu(max(1, int(max(0, float(w or 0)) / 100 * sw)))
    def NH(h):
        return Emu(max(1, int(max(0, float(h or 0)) / 100 * sh)))

    def crop_source(src_bytes, x, y, w, h):
        if not src_bytes:
            return None
        try:
            im = PILImage.open(io.BytesIO(src_bytes)).convert("RGB")
            iw, ih = im.size
            box = (int(x / 100 * iw), int(y / 100 * ih),
                   int((x + w) / 100 * iw), int((y + h) / 100 * ih))
            box = (max(0, box[0]), max(0, box[1]),
                   min(iw, max(box[0] + 1, box[2])),
                   min(ih, max(box[1] + 1, box[3])))
            if box[2] <= box[0] or box[3] <= box[1]:
                return None
            buf = io.BytesIO()
            im.crop(box).save(buf, format="PNG")
            return buf.getvalue()
        except Exception:
            return None

    n_slides = 0
    for idx, item in enumerate(slides[:80]):
        if not isinstance(item, dict):
            continue
        elements = item.get("elements")
        if not isinstance(elements, list) or not elements:
            continue
        src_bytes = sources[idx] if idx < len(sources) else None
        slide = prs.slides.add_slide(blank)
        n_slides += 1

        for el in elements[:_MAX_ELEMENTS]:
            if not isinstance(el, dict):
                continue
            etype = el.get("type")
            try:
                x, y = float(el.get("x", 0)), float(el.get("y", 0))
                w, h = float(el.get("w", 0)), float(el.get("h", 0))
            except (TypeError, ValueError):
                continue
            if w <= 0 or h <= 0:
                continue
            try:
                if etype == "text":
                    text = str(el.get("text", ""))
                    if not text.strip():
                        continue
                    txBox = slide.shapes.add_textbox(NX(x), NY(y), NW(w), NH(h))
                    tf = txBox.text_frame
                    tf.word_wrap = True
                    size = el.get("font_size", 18)
                    try:
                        size = max(6, min(72, float(size)))
                    except (TypeError, ValueError):
                        size = 18
                    r, g, b = _hex_to_rgb(el.get("color"))
                    align = _ALIGN_MAP.get(str(el.get("align", "left")).lower(),
                                           PP_ALIGN.LEFT)
                    first = True
                    for chunk in text.split("\n"):
                        p = tf.paragraphs[0] if first else tf.add_paragraph()
                        first = False
                        p.alignment = align
                        run = p.add_run()
                        run.text = chunk if chunk else " "
                        run.font.size = Pt(size)
                        run.font.bold = bool(el.get("bold", False))
                        run.font.italic = bool(el.get("italic", False))
                        run.font.color.rgb = RGBColor(r, g, b)
                        set_korean_font(run, font_name)
                elif etype == "shape":
                    shape_type = (MSO_SHAPE.OVAL if
                                  str(el.get("shape", "")).lower() == "ellipse"
                                  else MSO_SHAPE.RECTANGLE)
                    shp = slide.shapes.add_shape(shape_type, NX(x), NY(y),
                                                 NW(w), NH(h))
                    r, g, b = _hex_to_rgb(el.get("fill"), default=(0xF5, 0xF5, 0xF5))
                    shp.fill.solid()
                    shp.fill.fore_color.rgb = RGBColor(r, g, b)
                    shp.line.fill.background()  # 테두리 없음
                elif etype == "image":
                    png = crop_source(src_bytes, x, y, w, h)
                    if not png:
                        continue
                    slide.shapes.add_picture(io.BytesIO(png), NX(x), NY(y),
                                             NW(w), NH(h))
            except Exception:
                continue

        notes = str(item.get("notes", "")).strip()
        if notes:
            try:
                slide.notes_slide.placeholders[1].text = notes
            except Exception:
                pass

    if n_slides == 0:
        raise ConvertError("bad_json", "슬라이드 요소가 비어 있어요.")

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), {
        "mode": "precise",
        "pages": n_slides,
        "font": font_name,
        "aspect": aspect,
        "layout": True,
    }
